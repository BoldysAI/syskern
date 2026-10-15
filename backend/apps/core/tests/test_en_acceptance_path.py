"""Replayable English acceptance path (story 7.10).

Catalog consultation, catalog export, simulation export, then an English tariff
offer. A translated label must not stay the French source. Stored product text
(name, range) is not a label and stays as written.
"""

from __future__ import annotations

import re
from decimal import Decimal
from pathlib import Path

import pytest
from django.contrib.auth.models import User
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import Profile, Role
from apps.accounts.ui_locale import EN, FR, resolve_ui_locale
from apps.clients.models import Client
from apps.offers.services import excel as offer_excel
from apps.offers.tasks import generate_tariff_offers_task, offer_export_path
from apps.products.models import Product
from apps.products.tasks import export_products_task
from apps.simulations.models import Simulation, SimulationLine
from apps.simulations.tasks import export_simulation_task

pytestmark = pytest.mark.django_db

_BACKEND = Path(__file__).resolve().parents[3]
_REPO = _BACKEND.parent
_UI_ROOTS = (
    _REPO / "frontend" / "src" / "app" / "catalog",
    _REPO / "frontend" / "src" / "app" / "simulator",
    _REPO / "frontend" / "src" / "app" / "offers",
    _REPO / "frontend" / "src" / "components" / "layout" / "AppShell.tsx",
)
_API_FILES = (
    _BACKEND / "apps" / "products" / "views.py",
    _BACKEND / "apps" / "products" / "exports.py",
    _BACKEND / "apps" / "simulations" / "views.py",
    _BACKEND / "apps" / "simulations" / "exports.py",
    _BACKEND / "apps" / "offers" / "views.py",
)
_CATALOGS = _REPO / "frontend" / "src" / "lib" / "i18n" / "catalogs"
_ENTRY = re.compile(r'"([^"\\]+)"\s*:\s*"((?:\\.|[^"\\])*)"', re.DOTALL)
_UI_KEY = re.compile(r"""(?:\bt|translate)\(\s*["']([a-z][a-z0-9]*(?:\.[A-Za-z0-9_]+)+)["']""")
_API_KEY = re.compile(
    r"""translate(?:_request)?\(\s*(?:[^,\n]+,\s*)?["']([a-z][a-z0-9]*(?:\.[A-Za-z0-9_]+)+)["']"""
)


def _entries(source: str) -> dict[str, str]:
    return dict(_ENTRY.findall(source))


def _frontend_catalogs() -> tuple[dict[str, str], dict[str, str]]:
    french: dict[str, str] = {}
    english: dict[str, str] = {}
    for path in sorted(_CATALOGS.glob("*.ts")):
        french_src, english_src = path.read_text(encoding="utf-8").split("export const en", 1)
        french.update(_entries(french_src))
        english.update(_entries(english_src))
    return french, english


def _iter_sources(roots: tuple[Path, ...]) -> list[Path]:
    files: list[Path] = []
    for root in roots:
        if root.is_file():
            files.append(root)
            continue
        files.extend(path for path in root.rglob("*") if path.suffix in {".ts", ".tsx"})
    return files


def _ui_fallback_keys() -> list[str]:
    """Keys the path renders, whose English is missing so the screen stays French."""
    french, english = _frontend_catalogs()
    keys: set[str] = set()
    for path in _iter_sources(_UI_ROOTS):
        keys.update(_UI_KEY.findall(path.read_text(encoding="utf-8")))
    leaks: list[str] = []
    for key in sorted(keys):
        source = french.get(key)
        if source is None:
            leaks.append(f"{key} (unknown catalog key)")
            continue
        translated = english.get(key, "")
        if not translated.strip():
            leaks.append(f"{key} → {source}")
    return leaks


def _api_fallback_keys() -> list[str]:
    keys: set[str] = set()
    for path in _API_FILES:
        keys.update(_API_KEY.findall(path.read_text(encoding="utf-8")))
    leaks: list[str] = []
    for key in sorted(keys):
        if key not in FR:
            leaks.append(f"{key} (unknown catalog key)")
            continue
        translated = EN.get(key, "")
        if not translated.strip():
            leaks.append(f"{key} → {FR[key]}")
    return leaks


def _french_structural_labels() -> set[str]:
    """French copy that has a different English translation. Product text is absent."""
    banned = {french for key, french in FR.items() if EN.get(key, "") not in ("", french)}
    for spec in offer_excel._COLUMN_REGISTRY.values():
        french, english = spec["h"]["fr"], spec["h"]["en"]
        if french != english:
            banned.add(french)
    pairs = [offer_excel._SHEET_TITLE, offer_excel._TAGLINE, offer_excel._CONDITIONS_TEXT]
    pairs.extend(offer_excel._LABELS.values())
    for labels in pairs:
        if labels["fr"] != labels["en"]:
            banned.add(labels["fr"])
    return banned


def _structural_leaks(path: Path, banned: set[str]) -> list[str]:
    workbook = load_workbook(path)
    found: list[str] = []
    for sheet in workbook.worksheets:
        if sheet.title in banned:
            found.append(f"{path.name} sheet {sheet.title}")
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value in banned:
                    found.append(f"{path.name} {sheet.title}!{cell.coordinate} {cell.value}")
    return found


@pytest.mark.skipif(not _CATALOGS.is_dir(), reason="frontend/ is not mounted (backend container)")
def test_path_screens_do_not_fall_back_to_french() -> None:
    assert _ui_fallback_keys() == []


def test_path_api_messages_do_not_fall_back_to_french() -> None:
    assert _api_fallback_keys() == []


def test_catalog_to_english_offer_keeps_translated_labels_in_english(
    tmp_path, monkeypatch, settings
):
    settings.DEFAULT_UI_LOCALE = "fr"
    user = User.objects.create_user(
        username="shanghai@example.com",
        email="shanghai@example.com",
        password="GoodSecret123!",
    )
    Profile.objects.update_or_create(user=user, defaults={"role": Role.COMMERCIAL, "locale": "en"})
    user = User.objects.select_related("profile").get(pk=user.pk)
    assert resolve_ui_locale(user) == "en"

    product = Product.objects.create(
        sku_code="KCFF6A4PZHDBL5-21",
        name="Câble cat7",
        range="Catégorie 7",
        primary_packaging_qty=10,
        pallet_qty=9,
    )
    client = Client.objects.create(name="Shanghai desk")
    api = APIClient()
    api.force_authenticate(user)

    listed = api.get("/api/products/")
    assert listed.status_code == 200
    assert product.sku_code in listed.content.decode()

    detail = api.get(f"/api/products/{product.sku_code}/")
    assert detail.status_code == 200
    assert detail.json()["name"] == "Câble cat7"

    catalog_dir = tmp_path / "catalog"
    sim_dir = tmp_path / "sim"
    offer_dir = tmp_path / "offers"
    monkeypatch.setattr("apps.products.tasks.EXPORT_DIR", catalog_dir)
    monkeypatch.setattr("apps.simulations.tasks.EXPORT_DIR", sim_dir)
    monkeypatch.setattr("apps.offers.tasks.EXPORT_DIR", offer_dir)

    catalog_task = export_products_task.apply(
        kwargs={"ids": [str(product.id)], "locale": resolve_ui_locale(user)}
    )
    catalog_file = catalog_dir / f"{catalog_task.id}.xlsx"
    catalog_book = load_workbook(catalog_file)
    assert catalog_book.sheetnames == ["Product catalog"]
    assert catalog_book.active["A1"].value == "SKU"
    assert catalog_book.active["B1"].value == "Name"
    assert catalog_book.active["B2"].value == "Câble cat7"

    simulation = Simulation.objects.create(
        label="Tarif Q3 2026",
        simulation_type="tariff",
        status="draft",
        market_params={"fx_eur_usd": "1.15", "fx_eur_rmb": "7.95"},
    )
    SimulationLine.objects.create(
        simulation=simulation, product=product, pv_eur=Decimal("487.70"), status="ok"
    )
    Simulation.objects.filter(pk=simulation.pk).update(status="finalized")

    simulation_task = export_simulation_task.apply(
        args=[str(simulation.pk), resolve_ui_locale(user)]
    )
    simulation_file = sim_dir / f"{simulation_task.id}.xlsx"
    simulation_book = load_workbook(simulation_file)
    assert simulation_book.sheetnames[:2] == ["Summary", "Results"]

    generated = generate_tariff_offers_task.apply(
        args=[
            str(simulation.pk),
            {
                "client_ids": [str(client.id)],
                "columns": ["sku_code", "name", "unit_price"],
                "target_currency": "EUR",
                "language": "en",
                "expiration_date": "2026-09-30",
                "incoterm": "EXW",
                "label": "Tarif Q3 2026",
            },
        ]
    ).get()
    offer_file = offer_export_path(generated["offers"][0]["offer_id"])

    banned = _french_structural_labels()
    leaks = (
        _structural_leaks(catalog_file, banned)
        + _structural_leaks(simulation_file, banned)
        + _structural_leaks(offer_file, banned)
    )
    assert leaks == []
    offer_book = load_workbook(offer_file)
    assert "Tariffs" in offer_book.sheetnames
    flat = [
        cell.value for sheet in offer_book.worksheets for row in sheet.iter_rows() for cell in row
    ]
    assert "Unit price" in flat
    assert "Câble cat7" in flat
